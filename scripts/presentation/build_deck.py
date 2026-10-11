"""Собирает docs/scout-presentation.pptx из шаблона ЛЦТ 2026 ФСП.

Запуск из корня репозитория:
    python3 scripts/presentation/build_deck.py "ЛЦТ_2026_Шаблон презентации ФСП.pptx" docs/scout-presentation.pptx

Скрипт оставляет нужные слайды шаблона в заданном порядке и заполняет их.
Цифры валидации берутся из backend/eval/validation-report.json.
Места, которые команда заполняет сама, помечены квадратными скобками.
"""
import copy
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / 'frontend' / 'screenshots'
REPORT = json.loads((ROOT / 'backend' / 'eval' / 'validation-report.json').read_text())
TMP = Path(sys.argv[3]) if len(sys.argv) > 3 else Path('/tmp')

# Слайды шаблона (номер с 1) в порядке новой презентации.
ORDER = [7, 8, 9, 10, 11, 25, 27, 16, 20, 17, 22, 28, 24, 15, 14]

# Макеты шаблона красят пустые поля бледным цветом, поэтому цвет задаётся явно.
INK, DEEP, PINK, WHITE = '1C1D22', '520977', 'FF0053', 'FFFFFF'


def pct(x):
    return f'{x * 100:.1f}'.replace('.', ',') + '%'


def num(x, d=3):
    return f'{x:.{d}f}'.replace('.', ',')


A = REPORT['assessment']
M = REPORT['matching']['holdout']['systems']
LEAK = REPORT['bank_leak']


def by_ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(idx)


def by_id(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def set_text(shape, paras, size=None, bold=None, color=None, plain=False):
    """Заменяет текст, сохраняя оформление первого абзаца и первого фрагмента."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr = copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        if i and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        text, b, c = item if isinstance(item, tuple) and len(item) == 3 else \
            (*item, color) if isinstance(item, tuple) else (item, bold, color)
        r = p.add_run()
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        r.text = text
        if size:
            r.font.size = Pt(size)
        if b is not None:
            r.font.bold = b
        if c:
            r.font.color.rgb = RGBColor.from_string(c)
        if plain:
            ppr_el = p._p.get_or_add_pPr()
            ppr_el.set('marL', '0')
            ppr_el.set('indent', '0')
            for old in ppr_el.findall(qn('a:buNone')) + ppr_el.findall(qn('a:buChar')):
                ppr_el.remove(old)
            ppr_el.insert(0, ppr_el.makeelement(qn('a:buNone'), {}))


def fit_pill(slide, pill_id, title_shape, title):
    """Расширяет плашку заголовка под длину текста."""
    pill = by_id(slide, pill_id)
    w = Inches(0.8 + 0.17 * len(title))
    pill.width = max(pill.width, w)
    set_text(title_shape, title)


def shot(name, box=None, max_w=1600):
    """Скриншот в JPEG: обрезка и уменьшение, чтобы презентация не весила десятки мегабайт."""
    im = Image.open(SHOTS / name).convert('RGB')
    if box:
        im = im.crop(box)
    if im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
    path = TMP / (Path(name).stem + ('-crop' if box else '') + '.jpg')
    im.save(path, quality=85)
    return str(path)


def fill_ph(ph, path):
    """Вставляет картинку в плейсхолдер и сохраняет его место на слайде."""
    geom = ph.left, ph.top, ph.width, ph.height
    pic = ph.insert_picture(path)
    pic.left, pic.top, pic.width, pic.height = geom
    return pic


def remove(shape):
    shape._element.getparent().remove(shape._element)


def picture_in(slide, path, x, y, w, h):
    """Вписывает картинку в прямоугольник по центру, без искажения."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def main(template, out):
    prs = Presentation(template)
    ids = prs.slides._sldIdLst
    entries = list(ids)
    for e in entries:
        ids.remove(e)
    for n in ORDER:
        ids.append(entries[n - 1])
    s = list(prs.slides)

    # 1. Титул
    sl = s[0]
    set_text(by_ph(sl, 0), 'СКАУТ: найм по подтверждённым навыкам')
    set_text(by_ph(sl, 12), 'Цифровая платформа-агрегатор ИТ-вакансий с верифицированным профилем '
                            'достижений участника ФСП · команда [Название команды]')

    # 2. О команде и кратко о решении
    sl = s[1]
    hero = shot('result-scoreboard.png', (624, 96, 2688, 1056))
    fill_ph(by_ph(sl, 10), hero)
    set_text(by_ph(sl, 0), 'Скаут')
    set_text(by_id(sl, 5), 'Кандидат подтверждает навыки адаптивным тестом и получает категорию '
                           '«специализация · грейд». Работодатель ищет по подтверждённым категориям '
                           'и сам приглашает с вилкой.')
    set_text(by_id(sl, 8), 'Грейд ставит только тест, а не резюме. Задания генерируются из '
                           'параметрических семейств: слитые ответы почти не помогают.')

    # 3. Команда: заполняет команда
    set_text(by_ph(s[2], 0), 'Команда')

    # 4. История и сложности
    sl = s[3]
    set_text(by_ph(sl, 0), 'О команде и решении')
    set_text(by_id(sl, 40), 'Устойчивый банк при потолке экспозиции: решили параметрическими семействами. '
                            'Валидация нашла обрыв 48% сессий на редких компетенциях: исправили. '
                            'Задания с выбором угадывались по дампу банка: добавили пулы формулировок '
                            'с разными ответами и учёт угадывания (3PL).')

    # 5. Коротко о решении
    sl = s[4]
    set_text(by_ph(sl, 49), 'Техническая суть решения')
    set_text(by_ph(sl, 38), [
        'Адаптивный тест на IRT 3PL: грейд ставит тест, а не резюме',
        '244 параметрических семейства: у каждого кандидата свой экземпляр задания',
        'Практика с кодом в песочнице Piston',
        'Подбор по подтверждённым категориям с разложением балла на факторы',
        'Вход через ФСП ID (Keycloak), достижения поднимают внутри категории',
    ])
    set_text(by_ph(sl, 51), 'Ценность для ФСП и работодателей')
    set_text(by_ph(sl, 42), [
        'Работодатель сравнивает кандидатов по одной шкале, а не по самоописанию',
        'Кандидат один раз подтверждает навыки и получает приглашения с вилкой',
        'Контакты открываются только после согласия кандидата',
        'Участники ФСП получают канал трудоустройства, где видны их достижения',
    ])

    # 6. Путь: 5 шагов, номера чередуются сверху и снизу
    sl = s[5]
    fit_pill(sl, 12, by_ph(sl, 0), 'Как это работает')
    steps = [
        (26, 'Вход через ФСП ID', 'Согласия и опрос, достижения ФСП подтягиваются в профиль'),
        (32, 'Адаптивный тест', '14 заданий с практикой или экспресс из 6 заданий'),
        (28, 'Категория', '«Специализация · грейд» с перцентилем внутри категории'),
        (34, 'Подборка', 'Работодатель описывает потребность и видит объяснение балла'),
        (30, 'Приглашение', 'Вилка обязательна, контакты после согласия кандидата'),
    ]
    for idx, head, body in steps:
        set_text(by_ph(sl, idx), head, bold=True, color=DEEP, plain=True)
        set_text(by_ph(sl, idx + 1), body, color=INK, plain=True)

    # 7. Интерфейс в окнах браузера
    sl = s[6]
    fit_pill(sl, 2, by_ph(sl, 0), 'Интерфейс')
    fill_ph(by_ph(sl, 14), shot('result-scoreboard.png'))
    fill_ph(by_ph(sl, 18), shot('bank-player-cards.png'))
    set_text(by_ph(sl, 15), 'Кандидат: подтверждённая категория, перцентиль, навыки и правила пересдачи', color=INK, plain=True)
    set_text(by_ph(sl, 16), 'Работодатель: банк кандидатов с подтверждёнными навыками и ожиданиями по зарплате', color=INK, plain=True)
    for grp, tb in ((8, 10), (31, 33)):
        url = [x for x in by_id(sl, grp).shapes if x.shape_id == tb][0]
        set_text(url, 'scout · кандидат' if grp == 8 else 'scout · работодатель')

    # 8. Адаптивный тест: 4 карточки
    sl = s[7]
    fit_pill(sl, 10, by_ph(sl, 0), 'Адаптивный тест')
    cards = [
        (49, 37, 38, '3PL', 'Модель IRT',
         'Следующее задание выбирается по текущей оценке умения, с учётом угадывания'),
        (50, 39, 40, '244', 'Семейства заданий',
         '160 числовых, 70 с выбором ответа, 14 канареек. Ключи только на сервере'),
        (51, 41, 42, '0,4', 'Потолок экспозиции',
         'Семейство попадает не больше чем в 40% сессий, покрытие компетенций обязательно'),
        (52, 43, 44, 'Код', 'Практика',
         'Решение проверяется тестами в песочнице Piston. Провал практики не перекрыть баллом'),
    ]
    for big, head, body, b, h, t in cards:
        ph = by_ph(sl, big)
        ph.width = Inches(2.4)
        set_text(ph, b, size=24, bold=True, color=PINK, plain=True)
        set_text(by_ph(sl, head), h, bold=True, color=DEEP, plain=True)
        set_text(by_ph(sl, body), t, color=INK, plain=True)

    # 9. Подбор: фрагмент подборки и факторы
    sl = s[8]
    fit_pill(sl, 2, by_ph(sl, 0), 'Подбор с объяснением')
    remove(by_ph(sl, 14))
    league = shot('search-league-table.png', (640, 140, 1790, 1180))
    box = by_id(sl, 15)
    pad = Inches(0.15)
    picture_in(sl, league, box.left + pad, box.top + pad, box.width - 2 * pad, box.height - 2 * pad)
    factors = [
        'Подтверждённая категория и перцентиль: главный сигнал',
        'Навыки, подтверждённые тестом, весят больше заявленных',
        'Ожидания в вилке, отрасль и формат работы',
        'Регулярные задания и достижения ФСП: бонус внутри категории',
        'Грейд ниже запроса виден с пометкой и штрафом, а не скрыт',
    ]
    for i, f in enumerate(factors):
        set_text(by_ph(sl, 15 + i), f, color=INK)

    # 10. Валидация теста: 3 крупные цифры
    sl = s[9]
    set_text(by_ph(sl, 0), 'Валидация теста на 2000 кандидатах', color=DEEP)
    g = A['grade_by_theta']
    fx = A['fixed_12_item_form']
    stats = [
        (49, 37, 38, pct(g['within_one']['value']), 'Грейд ±1 ступень',
         f"Точно из 5 классов {pct(g['exact']['value'])}, κ {num(g['quadratic_kappa']['value'], 2)}. "
         f"Фиксированная форма из 12 заданий: {pct(fx['exact']['value'])}"),
        (50, 39, 40, '+' + num(LEAK['cat']['key_100']['inflation_vs_clean']['value'], 2),
         'Завышение при сливе всех ключей',
         f"У фиксированного теста +{num(LEAK['fixed_12']['leak_100']['inflation_vs_clean']['value'], 2)}. "
         'Каждый кандидат получает свой экземпляр задания'),
        (51, 41, 42, num(REPORT['test_retest']['quadratic_kappa']['value'], 2), 'κ повторного теста',
         'Вторая попытка с другими заданиями. Ложные отказы '
         f"{pct(A['confirmation_decision']['false_reject_rate']['value'])}, ложные подтверждения "
         f"{pct(A['confirmation_decision']['false_confirm_rate']['value'])}"),
    ]
    for big, head, body, b, h, t in stats:
        ph = by_ph(sl, big)
        ph.width = Inches(3.0)
        set_text(ph, b, size=32, bold=True, color=PINK, plain=True)
        set_text(by_ph(sl, head), h, bold=True, color=DEEP, plain=True)
        set_text(by_ph(sl, body), t, color=INK, plain=True)

    # 11. Валидация подбора: диаграмма nDCG@10
    sl = s[10]
    fit_pill(sl, 2, by_ph(sl, 0), 'Валидация подбора')
    chart = [x for x in sl.shapes if x.has_chart][0].chart
    cd = CategoryChartData(number_format='0.000')
    cd.categories = ['BM25 по резюме', 'BM25 + фильтры', 'Скаут без теста', 'Скаут']
    cd.add_series('nDCG@10', [M['bm25']['ndcg_at_10']['value'], M['bm25_filtered']['ndcg_at_10']['value'],
                              M['ours_no_test']['ndcg_at_10']['value'], M['ours']['ndcg_at_10']['value']])
    chart.replace_data(cd)
    chart.value_axis.maximum_scale = 0.7
    chart.value_axis.minimum_scale = 0
    chart.value_axis.major_unit = 0.1
    chart.value_axis.tick_labels.number_format = '0.0'
    chart.value_axis.tick_labels.number_format_is_linked = False
    plot = chart.plots[0]
    for point, value in zip(plot.series[0].points, plot.series[0].values):
        tf = point.data_label.text_frame
        tf.text = num(value)
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(INK)
        point.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    up = M['ours']['ndcg_at_10']['value'] / M['bm25']['ndcg_at_10']['value'] - 1
    pairs = [
        (21, 18, f'+{up * 100:.0f}% к BM25',
         f"nDCG@10 {num(M['ours']['ndcg_at_10']['value'])} против {num(M['bm25']['ndcg_at_10']['value'])} "
         'на отложенных 210 потребностях'),
        (22, 23, f"P@10 {num(M['ours']['p_at_10']['value'])}",
         f"Две трети первой десятки релевантны, у BM25 {num(M['bm25']['p_at_10']['value'])}"),
        (24, 25, f"+{num(M['ours']['ndcg_at_10']['value'] - M['ours_no_test']['ndcg_at_10']['value'])} от теста",
         'Без результатов теста та же формула теряет четверть качества'),
        (26, 27, 'Веса подобраны на dev', 'Проверены на выборке, которая в подборе не участвовала. '
                                          'Всё воспроизводит npm run eval'),
    ]
    for head, body, h, t in pairs:
        set_text(by_ph(sl, head), h, bold=True, color=PINK, plain=True)
        set_text(by_ph(sl, body), t, color=INK, plain=True)

    # 12. Мобильная версия
    sl = s[11]
    fit_pill(sl, 3, by_ph(sl, 0), 'Мобильная версия')
    fill_ph(by_ph(sl, 14), shot('mobile-search.png'))
    calls = [
        (21, 18, '375 px', '30 экранов без горизонтальной прокрутки'),
        (22, 23, 'Работодатель', 'Подборка, карточка кандидата и приглашение с телефона'),
        (24, 25, 'Кандидат', 'Тест, инвайты и профиль в мобильном браузере'),
        (26, 27, 'Две темы', 'Светлая и тёмная тема во всех кабинетах'),
    ]
    for head, body, h, t in calls:
        set_text(by_ph(sl, head), h, bold=True, color=DEEP, plain=True)
        set_text(by_ph(sl, body), t, color=INK, plain=True)

    # 13. Архитектура
    sl = s[12]
    fit_pill(sl, 2, by_ph(sl, 0), 'Архитектура и запуск')
    arch = [
        (26, 'Интерфейс', ['React, Vite, TypeScript', 'nginx или Caddy с HTTPS', 'Контракт OpenAPI, клиент генерируется']),
        (31, 'Сервер', ['NestJS, Prisma, PostgreSQL', 'Keycloak: OIDC с PKCE', 'Redis и BullMQ, Piston для кода']),
        (32, 'Запуск и качество', ['docker compose up одной командой', '420 демо-кандидатов, учётки жюри',
                                   'CI: 130 тестов бэкенда, Playwright, eval']),
    ]
    for idx, head, lines in arch:
        set_text(by_ph(sl, idx), [(head, True, DEEP)] + [(x, False, INK) for x in lines])

    # 14. Развитие
    sl = s[13]
    set_text(by_ph(sl, 0), 'Развитие продукта', color=DEEP)
    plans = [
        'Пилот на реальных кандидатах и калибровка параметров заданий по их ответам',
        'LLM (GigaChat, резерв YandexGPT): черновики заданий с проверкой экспертом',
        'Боевой реестр ФСП и ФСП ID вместо эмулятора',
        'Новые специализации и практические кейсы в банке заданий',
        'Выгрузка приглашённых кандидатов в ATS работодателя',
    ]
    for i, t in enumerate(plans):
        ph = by_ph(sl, 15 + i)
        ph.width = Inches(12.2)
        set_text(ph, t, color=INK)

    # 15. Демонстрация
    sl = s[14]
    set_text(by_ph(sl, 0), 'Демонстрация', color=DEEP)
    set_text(by_ph(sl, 14), [
        ('Репозиторий', True, DEEP), ('github.com/Egor2809/fsp-scout', False, INK),
        ('Стенд и видео', True, DEEP), ('[ссылка на стенд] · [ссылка на видео]', False, INK),
        ('Экспресс-тест', True, DEEP), ('Кнопка при старте теста: 6 заданий за 3–5 минут', False, INK),
        ('Учётки жюри', True, DEEP), ('jury1…5-candidate@demo.ru и jury1…5-employer@demo.ru, пароль demo12345', False, INK),
    ])
    remove(by_ph(sl, 10))
    picture_in(sl, shot('mobile-candidate-card.png'), Inches(7.2), Inches(1.3), Inches(5.6), Inches(5.6))

    prs.save(out)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
